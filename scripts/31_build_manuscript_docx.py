#!/usr/bin/env python3
"""Build the editable manuscript DOCX from the version-controlled Markdown.

This is a presentation-only build. It reads the manuscript and already
committed figures; it does not calculate or alter scientific results.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / "manuscript" / "manuscript.md"
DEFAULT_OUTPUT = ROOT / "manuscript" / "manuscript.docx"

NAVY = "17365D"
PALE_BLUE = "EAF1F8"
PALE_GRAY = "F7F7F7"
GRID_GRAY = "D9D9D9"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build the manuscript DOCX without recomputing scientific results."
    )
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def set_run_font(run, name: str, size: float | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 90, start: int = 90, bottom: int = 90, end: int = 90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = GRID_GRAY, size: int = 5) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instruction, separate, text, end):
        run._r.append(node)
    set_run_font(run, "Times New Roman", 9)


INLINE_PATTERN = re.compile(r"(`[^`]+`|\*[^*]+\*)")


def add_inline(paragraph, text: str, *, size: float = 11.5) -> None:
    for part in INLINE_PATTERN.split(text):
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, "Courier New", max(9.0, size - 1.0))
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
            set_run_font(run, "Times New Roman", size)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, "Times New Roman", size)


def markdown_table(lines: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if cells and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            continue
        rows.append(cells)
    return rows


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    n_cols = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.style = "Table Grid"
    set_table_borders(table)
    for row_index, values in enumerate(rows):
        tr_pr = table.rows[row_index]._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        if row_index == 0:
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            tr_pr.append(header)
        for col_index in range(n_cols):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            value = values[col_index] if col_index < len(values) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            # Keep the final two rows together to avoid a one-row table orphan,
            # while still allowing long tables to use remaining page space.
            paragraph.paragraph_format.keep_with_next = row_index == len(rows) - 2
            is_numeric = bool(re.fullmatch(r"[-+−]?\d[\d.,%/\- ]*", value))
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 or is_numeric else WD_ALIGN_PARAGRAPH.LEFT
            )
            add_inline(paragraph, value, size=8.5)
            for run in paragraph.runs:
                if row_index == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
            if row_index == 0:
                set_cell_shading(cell, NAVY)
            elif row_index % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
            else:
                set_cell_shading(cell, "FFFFFF")
    after = document.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def image_size(path: Path, max_width: float = 6.45, max_height: float = 8.05) -> tuple[float, float]:
    from PIL import Image

    with Image.open(path) as image:
        width_px, height_px = image.size
    ratio = width_px / height_px
    width = max_width
    height = width / ratio
    if height > max_height:
        height = max_height
        width = height * ratio
    return width, height


def set_image_alt_text(run, description: str) -> None:
    drawing = run._r.find(qn("w:drawing"))
    if drawing is None:
        return
    doc_pr = drawing.find(".//" + qn("wp:docPr"))
    if doc_pr is not None:
        doc_pr.set("descr", description)


def add_figure(document: Document, source: Path, description: str) -> None:
    if not source.exists():
        raise FileNotFoundError(f"Manuscript figure is missing: {source}")
    width, _height = image_size(source)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run()
    run.add_picture(str(source), width=Inches(width))
    set_image_alt_text(run, description)
    caption = document.add_paragraph()
    caption.style = document.styles["Caption"]
    caption.alignment = WD_ALIGN_PARAGRAPH.LEFT
    caption.paragraph_format.keep_together = True
    caption.paragraph_format.space_after = Pt(10)
    add_inline(caption, description, size=9.5)


def configure_document(document: Document, title: str) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(11.5)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True

    title_style = document.styles["Title"]
    title_style.font.name = "Times New Roman"
    title_style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    title_style.font.size = Pt(18)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.paragraph_format.space_after = Pt(10)

    for style_name, size, before, after in (
        ("Heading 1", 14, 12, 6),
        ("Heading 2", 12, 10, 4),
    ):
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = document.styles["Caption"]
    caption.font.name = "Times New Roman"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    caption.font.size = Pt(9.5)
    caption.font.italic = False
    caption.font.color.rgb = RGBColor(0, 0, 0)

    add_page_number(section.footer.paragraphs[0])

    document.core_properties.title = title
    document.core_properties.subject = "Lake Erken Sentinel-2 temporal reconstruction"
    document.core_properties.creator = ""
    document.core_properties.last_modified_by = ""


def build(source: Path, output: Path) -> Path:
    source = source.resolve()
    output = output.resolve()
    lines = source.read_text(encoding="utf-8").splitlines()
    title_line = next((line for line in lines if line.startswith("# ")), None)
    if title_line is None:
        raise ValueError("Manuscript source requires a level-one title.")
    title = title_line[2:].strip()

    document = Document()
    configure_document(document, title)
    index = 0
    first_title = True
    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        if not stripped:
            index += 1
            continue
        if stripped.startswith("|"):
            table_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            add_table(document, markdown_table(table_lines))
            continue
        image_match = re.fullmatch(r"!\[(.+)]\((.+)\)", stripped)
        if image_match:
            description, raw_path = image_match.groups()
            figure_path = (source.parent / raw_path).resolve()
            add_figure(document, figure_path, description)
            index += 1
            continue
        if stripped.startswith("# "):
            if first_title:
                paragraph = document.add_paragraph(style="Title")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(paragraph, stripped[2:].strip(), size=18)
                first_title = False
            index += 1
            continue
        if stripped.startswith("## "):
            document.add_paragraph(stripped[3:].strip(), style="Heading 1")
            index += 1
            continue
        if stripped.startswith("### "):
            document.add_paragraph(stripped[4:].strip(), style="Heading 2")
            index += 1
            continue
        if stripped.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(2)
            add_inline(paragraph, stripped[2:].strip())
            index += 1
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            next_line = lines[index].strip()
            if not next_line or next_line.startswith(("#", "|", "![", "- ")):
                break
            paragraph_lines.append(next_line)
            index += 1
        paragraph_text = " ".join(paragraph_lines)
        paragraph = document.add_paragraph()
        if paragraph_text.startswith("Running title:"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(12)
            run = paragraph.add_run(paragraph_text)
            run.italic = True
            set_run_font(run, "Times New Roman", 10.5)
        elif paragraph_text.startswith("Keywords:"):
            paragraph.paragraph_format.space_after = Pt(10)
            label, value = paragraph_text.split(":", 1)
            run = paragraph.add_run(label + ":")
            run.bold = True
            set_run_font(run, "Times New Roman", 10.5)
            add_inline(paragraph, value, size=10.5)
        else:
            add_inline(paragraph, paragraph_text)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return output


def main() -> int:
    args = parse_args()
    destination = build(args.source, args.output)
    print(destination)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
